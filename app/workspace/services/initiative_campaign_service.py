from io import BytesIO
import random
from typing import Any
from datetime import datetime
from zoneinfo import ZoneInfo

from flask import current_app
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.database.transaction import transactional
from app.workspace.repositories.initiative_campaign_repository import (
    InitiativeCampaignRepository,
)
from app.workspace.repositories.initiative_repository import InitiativeRepository
from app.workspace.repositories.linear_motion_repository import LinearMotionRepository


TRACKING_HEADERS = [
    "_ID", "Cliente", "Ciudad", "Prioridad", "Estado", "Fecha de contacto",
    "Contacto en cliente", "Aplicación o máquina", "Marca o referencia instalada",
    "Interés", "Próxima acción", "Fecha compromiso", "¿Oportunidad?",
]

INSIDE_SALES_REPS = (
    ("RICARDO CASTANEDA", "ricardo.castaneda@lugohermanos.com"),
    ("HAROL CAMARGO", "harol.camargo@lugohermanos.com"),
    ("TOMAS VARGAS", "asesor1@lugohermanos.com"),
    ("FERNEY RINCON", "ferney.rincon@lugohermanos.com"),
    ("DUBAN QUINTERO", "duban.quintero@lugohermanos.com"),
)


class InitiativeCampaignService:

    @staticmethod
    def run_followup_automation(now: datetime | None = None) -> dict[str, int]:
        """Send two quiet reminders, then expose inactive advisors to management."""
        if current_app.config.get("EMAIL_PROVIDER") != "microsoft":
            return {"reminders": 0, "escalations": 0}
        local_now = now or datetime.now(ZoneInfo("America/Bogota"))
        if local_now.weekday() >= 5 or local_now.hour < 8:
            return {"reminders": 0, "escalations": 0}
        reminders = escalations = 0
        provider = current_app.extensions["email_provider"]
        for record in InitiativeCampaignRepository.list_files():
            if not record.get("sent_at") or not record.get("recipient"):
                continue
            if InitiativeCampaignRepository.advisor_has_progress(
                record["initiative_id"], record["advisor"],
            ):
                if record.get("escalated_at"):
                    InitiativeCampaignRepository.resolve_escalation(record["id"])
                continue
            sent = datetime.fromisoformat(record["sent_at"]).replace(
                tzinfo=local_now.tzinfo,
            )
            business_days = InitiativeCampaignService._business_days(sent, local_now)
            if business_days >= 3 and not record.get("first_reminder_at"):
                InitiativeCampaignService._send_reminder(provider, record, 1)
                InitiativeCampaignRepository.mark_automation(
                    record["id"], first=True, note="Primer recordatorio enviado",
                )
                reminders += 1
            elif business_days >= 7 and not record.get("second_reminder_at"):
                InitiativeCampaignService._send_reminder(provider, record, 2)
                InitiativeCampaignRepository.mark_automation(
                    record["id"], second=True, note="Segundo recordatorio enviado",
                )
                reminders += 1
            elif record.get("second_reminder_at") and not record.get("escalated_at"):
                second = datetime.fromisoformat(record["second_reminder_at"]).replace(
                    tzinfo=local_now.tzinfo,
                )
                if InitiativeCampaignService._business_days(second, local_now) >= 2:
                    InitiativeCampaignRepository.mark_automation(
                        record["id"], escalated=True,
                        note="Sin actualización después del segundo recordatorio",
                    )
                    InitiativeRepository.create_event(
                        initiative_id=record["initiative_id"],
                        event_type="campaign_escalation",
                        title="Seguimiento escalado a gerencia",
                        details=(f"{record['advisor']} no registró avances después "
                                 "de dos recordatorios automáticos."),
                        created_by="automation",
                    )
                    escalations += 1
        return {"reminders": reminders, "escalations": escalations}

    @staticmethod
    def _business_days(start: datetime, end: datetime) -> int:
        day = start.date()
        total = 0
        while day < end.date():
            day = day.fromordinal(day.toordinal() + 1)
            if day.weekday() < 5:
                total += 1
        return total

    @staticmethod
    def _send_reminder(provider, record: dict[str, Any], number: int) -> None:
        ordinal = "primer" if number == 1 else "segundo y último"
        subject = f"Recordatorio {number} de 2 – Campaña THK plásticos"
        link = record["web_url"]
        provider.send(
            sender=current_app.config.get("MICROSOFT_MAILBOX_ADDRESS") or "",
            recipients=[record["recipient"]], cc=[], subject=subject,
            body_text=(f"Este es el {ordinal} recordatorio para actualizar el archivo "
                       f"de seguimiento THK: {link}"),
            body_html=f"""
            <p>Estimado(a) {record['advisor'].title()}:</p>
            <p>Este es el <strong>{ordinal} recordatorio</strong> para registrar el
            avance de los clientes asignados en la iniciativa THK para plásticos.</p>
            <p>Por favor, actualice directamente en Excel Online el estado del
            contacto y la información disponible. Si todavía no ha logrado
            comunicación, seleccione <strong>Contacto intentado</strong>.</p>
            <p><a href="{link}"><strong>Abrir archivo de seguimiento</strong></a></p>
            <p>No es necesario descargar ni enviar el archivo por correo.</p>
            <p>Cordialmente,<br><strong>Lugo Hermanos</strong></p>
            """,
        )

    @staticmethod
    def sync_all_from_onedrive() -> dict[str, int]:
        """Synchronize linked campaigns without one failure stopping the rest."""
        if current_app.config.get("EMAIL_PROVIDER") != "microsoft":
            return {"processed": 0, "succeeded": 0, "failed": 0}

        files = InitiativeCampaignRepository.list_files()
        succeeded = 0
        failed = 0
        for file_record in files:
            initiative_id = file_record["initiative_id"]
            advisor = file_record["advisor"]
            try:
                InitiativeCampaignService.sync_from_onedrive(
                    initiative_id, advisor=advisor,
                )
                succeeded += 1
            except Exception as exc:
                failed += 1
                message = str(exc).strip() or exc.__class__.__name__
                current_app.logger.exception(
                    "No se pudo sincronizar la campaña de la iniciativa %s",
                    initiative_id,
                )
                InitiativeCampaignRepository.mark_sync(
                    initiative_id,
                    advisor=advisor,
                    etag=None,
                    status="error",
                    message=message[:500],
                )
        return {
            "processed": len(files),
            "succeeded": succeeded,
            "failed": failed,
        }

    @staticmethod
    def get_campaign(initiative_id: int) -> dict[str, Any]:
        customers = InitiativeCampaignRepository.list_customers(initiative_id)
        total = len(customers)
        contacted = sum(c["status"] not in {"Pendiente", "Contacto intentado"} for c in customers)
        applications = sum(bool(c.get("application")) for c in customers)
        opportunities = sum(c.get("opportunity_decision") == "Sí" for c in customers)
        inside_sales = []
        for advisor, email in INSIDE_SALES_REPS:
            assigned = sum(c["advisor"] == advisor for c in customers)
            if assigned:
                inside_sales.append({
                    "advisor": advisor,
                    "email": email,
                    "assigned": assigned,
                })
        return {
            "customers": customers,
            "file": InitiativeCampaignRepository.get_file(initiative_id),
            "files": [
                item for item in InitiativeCampaignRepository.list_files()
                if item["initiative_id"] == initiative_id
            ],
            "escalations": [
                item for item in InitiativeCampaignRepository.list_files()
                if item["initiative_id"] == initiative_id and item.get("escalated_at")
            ],
            "inside_sales": inside_sales,
            "summary": {
                "assigned": total,
                "contacted": contacted,
                "contacted_pct": round(contacted * 100 / total) if total else 0,
                "applications": applications,
                "applications_pct": round(applications * 100 / total) if total else 0,
                "opportunities": opportunities,
                "opportunities_pct": round(opportunities * 100 / total) if total else 0,
            },
        }

    @staticmethod
    @transactional
    def distribute_warehouse_customers(initiative_id: int) -> dict[str, Any]:
        """Randomly balance ERP warehouse accounts across inside sales."""
        opportunity = LinearMotionRepository.segment_opportunity("PLASTICOS", "THK")
        warehouse_customers = [
            item for item in opportunity.get("customers", [])
            if str(item.get("advisor") or "").strip().casefold().startswith("almacen")
        ]
        if not warehouse_customers:
            raise ValueError("No se encontraron clientes de Almacén para distribuir.")

        existing = {
            str(item["customer_erp_id"]): item
            for item in InitiativeCampaignRepository.list_customers(initiative_id)
        }
        rep_names = [name for name, _ in INSIDE_SALES_REPS]
        counts = {name: 0 for name in rep_names}
        pending = []
        for customer in warehouse_customers:
            customer_id = str(customer.get("customer_id") or "").strip()
            assigned = existing.get(customer_id, {}).get("advisor")
            if assigned in counts:
                counts[assigned] += 1
            else:
                pending.append(customer)

        rng = random.SystemRandom()
        rng.shuffle(pending)
        for customer in pending:
            minimum = min(counts.values())
            eligible = [name for name, count in counts.items() if count == minimum]
            advisor = rng.choice(eligible)
            customer_id = str(customer.get("customer_id") or "").strip()
            InitiativeCampaignRepository.upsert_customer(initiative_id, {
                "customer_erp_id": customer_id,
                "customer_name": customer.get("name") or customer_id,
                "city": customer.get("city"),
                "advisor": advisor,
                "priority": "Alta" if counts[advisor] < 5 else "Media",
                "brands": customer.get("brands"),
                "competitor_linear_brands": customer.get("competitor_linear_brands"),
                "last_purchase": customer.get("last_purchase"),
                "other_sales": customer.get("other_sales"),
                "status": "Pendiente",
                "source_key": f"THK-PLASTICOS-{customer_id}",
            })
            counts[advisor] += 1

        InitiativeRepository.create_event(
            initiative_id=initiative_id,
            event_type="campaign_assignment",
            title="Clientes de Almacén distribuidos a Inside Sales",
            details=(f"{len(warehouse_customers)} clientes distribuidos aleatoriamente "
                     "y de forma equitativa entre cinco asesores."),
            created_by="automation",
        )
        return {
            "total": len(warehouse_customers),
            "assignments": [
                {"advisor": name, "email": email, "assigned": counts[name]}
                for name, email in INSIDE_SALES_REPS
            ],
        }

    @staticmethod
    @transactional
    def launch_microsoft_pilot(
        initiative_id: int,
        *, recipient: str,
    ) -> dict[str, Any]:
        customers = InitiativeCampaignRepository.list_customers(initiative_id)
        if not customers:
            raise ValueError("La iniciativa no tiene clientes de campaña.")
        advisor = customers[0]["advisor"]
        return InitiativeCampaignService.launch_advisor_campaign.__wrapped__(
            initiative_id, advisor=advisor, recipient=recipient,
        )

    @staticmethod
    @transactional
    def launch_advisor_campaign(
        initiative_id: int, *, advisor: str, recipient: str,
        prepare_if_missing: bool = False, reset_tracking: bool = True,
    ) -> dict[str, Any]:
        if prepare_if_missing:
            InitiativeCampaignService._prepare_advisor_customers(
                initiative_id, advisor,
            )
        customers = [
            item for item in InitiativeCampaignRepository.list_customers(initiative_id)
            if item["advisor"] == advisor
        ]
        if not customers:
            raise ValueError("El asesor no tiene clientes preparados para la campaña.")
        if reset_tracking:
            InitiativeCampaignRepository.reset_tracking(initiative_id, advisor)
        stream, filename = InitiativeCampaignService.export_workbook(
            initiative_id, advisor=advisor,
        )
        provider = current_app.extensions["email_provider"]
        if current_app.config.get("EMAIL_PROVIDER") != "microsoft":
            raise ValueError("El proveedor activo debe ser Microsoft 365.")
        uploaded = provider.upload_onedrive_file(
            filename=filename,
            content=stream.getvalue(),
            recipient=recipient,
        )
        InitiativeCampaignRepository.save_file(
            initiative_id, advisor, recipient, uploaded,
        )
        link = uploaded["web_url"]
        body_html = InitiativeCampaignService._campaign_email_html(
            advisor, len(customers), link,
        )
        provider.send(
            sender=current_app.config.get("MICROSOFT_MAILBOX_ADDRESS") or "",
            recipients=[recipient],
            cc=[],
            subject="Iniciativa comercial THK – Clientes asignados del segmento de plásticos",
            body_text=(
                f"Estimado(a) {advisor.title()}, consulte su archivo privado "
                f"de seguimiento en Excel Online: {link}"
            ),
            body_html=body_html,
        )
        InitiativeRepository.create_event(
            initiative_id=initiative_id,
            event_type="campaign_launch",
            title="Campaña THK enviada",
            details=(f"Archivo privado de {advisor} creado y correo enviado "
                     f"a {recipient} con {len(customers)} clientes."),
            created_by="microsoft",
        )
        return uploaded

    @staticmethod
    @transactional
    def sync_from_onedrive(
        initiative_id: int, advisor: str | None = None,
    ) -> dict[str, int]:
        file_record = InitiativeCampaignRepository.get_file(
            initiative_id, advisor,
        )
        if not file_record:
            raise ValueError("La iniciativa todavía no tiene archivo en OneDrive.")
        provider = current_app.extensions["email_provider"]
        content, etag = provider.download_onedrive_file(file_record["item_id"])
        result = InitiativeCampaignService.import_workbook.__wrapped__(
            initiative_id, BytesIO(content)
        )
        InitiativeCampaignRepository.mark_sync(
            initiative_id,
            advisor=file_record["advisor"],
            etag=etag,
            status="success",
            message=f"{result['updated']} clientes actualizados.",
        )
        return result

    @staticmethod
    def _prepare_advisor_customers(initiative_id: int, advisor: str) -> None:
        existing = [
            item for item in InitiativeCampaignRepository.list_customers(initiative_id)
            if item["advisor"] == advisor
        ]
        if existing:
            return
        opportunity = LinearMotionRepository.segment_opportunity("PLASTICOS", "THK")
        selected = [
            item for item in opportunity.get("customers", [])
            if str(item.get("advisor") or "").strip().casefold()
            == advisor.strip().casefold()
        ]
        for index, item in enumerate(selected):
            customer_id = str(item.get("customer_id") or "").strip()
            InitiativeCampaignRepository.upsert_customer(initiative_id, {
                "customer_erp_id": customer_id,
                "customer_name": item.get("name") or customer_id,
                "city": item.get("city"),
                "advisor": advisor,
                "priority": "Alta" if index < 5 else "Media",
                "brands": item.get("brands"),
                "competitor_linear_brands": item.get("competitor_linear_brands"),
                "last_purchase": item.get("last_purchase"),
                "other_sales": item.get("other_sales"),
                "status": "Pendiente",
                "source_key": f"THK-PLASTICOS-{customer_id}",
            })

    @staticmethod
    def _campaign_email_html(advisor: str, count: int, link: str) -> str:
        return f"""
        <p>Estimado(a) {advisor.title()}:</p>
        <p>Como parte de la iniciativa comercial para el desarrollo de la marca
        THK, se ha preparado una lista de <strong>{count} clientes activos del
        segmento de plásticos</strong> que actualmente no registran compras de
        esta marca.</p>
        <p>El objetivo es contactar a estos clientes e identificar máquinas o
        aplicaciones con movimiento lineal en las que exista la posibilidad de
        aportar una solución y desarrollar una eventual homologación.</p>
        <p>En esta primera etapa no se busca ofrecer una referencia específica.
        Se requiere conocer la máquina o aplicación, el componente instalado,
        la necesidad del cliente, su interés en evaluar THK y la siguiente acción.</p>
        <p><a href="{link}"><strong>Abrir archivo privado en Excel Online</strong></a></p>
        <p>El archivo contiene exclusivamente los clientes asignados para esta
        iniciativa y tiene tres pestañas:</p>
        <h3>1. Seguimiento</h3>
        <p>Es la única pestaña que debe actualizarse. Registre el estado, fecha y
        persona de contacto, aplicación o máquina, marca o referencia instalada,
        interés, próxima acción, fecha compromiso y si existe una oportunidad.</p>
        <p>Estados disponibles: <strong>Pendiente</strong>, <strong>Contacto
        intentado</strong>, <strong>Contactado</strong>, <strong>Aplicación
        identificada</strong> y <strong>Sin oportunidad</strong>. Marque
        “Sí” en oportunidad solamente cuando exista una necesidad concreta,
        interés y una siguiente acción comercial o técnica.</p>
        <h3>2. Contexto</h3>
        <p>Es únicamente de consulta y no debe modificarse. Presenta la última
        compra, compras de otras marcas, marcas históricas y posibles marcas
        relacionadas con movimiento lineal.</p>
        <h3>3. Instrucciones</h3>
        <p>Resume el objetivo, las fechas y la información que debe recopilarse.</p>
        <h3>Apoyo para homologaciones</h3>
        <p>Si el cliente utiliza otra marca y desea evaluar THK, registre la
        aplicación, marca, referencia disponible y necesidad, y seleccione
        <strong>Proponer homologación</strong>. Lugo Hermanos brindará el apoyo
        comercial y técnico para revisar la equivalencia y acompañar el proceso.</p>
        <h3>Manejo del archivo</h3>
        <p>Trabaje directamente en Excel Online. Los cambios se guardan
        automáticamente. No es necesario descargar, copiar ni devolver el archivo.</p>
        <p><strong>Fechas de seguimiento:</strong><br>
        Primera actualización: 27 de octubre de 2026.<br>
        Segundo y último plazo: 2 de noviembre de 2026.<br>
        Los casos sin actualización se escalarán a Gerencia a partir del
        4 de noviembre de 2026.</p>
        <p>Cordialmente,<br><strong>Lugo Hermanos</strong><br>
        Equipo comercial y técnico</p>
        """

    @staticmethod
    def export_workbook(initiative_id: int, advisor: str | None = None) -> tuple[BytesIO, str]:
        initiative = InitiativeRepository.get_initiative(initiative_id)
        if not initiative:
            raise ValueError("La iniciativa no existe.")
        customers = InitiativeCampaignRepository.list_customers(initiative_id)
        if advisor:
            customers = [item for item in customers if item["advisor"] == advisor]
        if not customers:
            raise ValueError("No hay clientes de campaña para exportar.")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Seguimiento"
        sheet.append([initiative["name"]])
        sheet.append(["Actualice las columnas azules. El Command Center usa el ID para sincronizar cada cliente."])
        sheet.append(TRACKING_HEADERS)
        for customer in customers:
            sheet.append([
                customer["source_key"], customer["customer_name"], customer.get("city"),
                customer.get("priority"), customer.get("status"), customer.get("contact_date"),
                customer.get("contact_name"), customer.get("application"),
                customer.get("installed_reference"), customer.get("interest"),
                customer.get("next_action"), customer.get("next_action_date"),
                customer.get("opportunity_decision"),
            ])

        dark_blue, light_blue, orange = "17365D", "DDEBF7", "F4B183"
        sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor=dark_blue)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=13)
        for cell in sheet[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=dark_blue)
            cell.alignment = Alignment(wrap_text=True)
        for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=5, max_col=13):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=light_blue)
        for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=1, max_col=4):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
        sheet.column_dimensions["A"].hidden = True
        widths = [3, 34, 18, 12, 24, 18, 24, 34, 30, 13, 24, 18, 16]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        sheet.freeze_panes = "E4"
        sheet.auto_filter.ref = f"B3:M{sheet.max_row}"

        validations = {
            "E": '"Pendiente,Contacto intentado,Contactado,Aplicación identificada,Sin oportunidad"',
            "J": '"Alto,Medio,Bajo,Ninguno"',
            "K": '"Llamar,Enviar información,Visitar,Levantar datos,Cotizar,Proponer homologación"',
            "M": '"Sí,No,Revisar"',
        }
        for column, formula in validations.items():
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            sheet.add_data_validation(validation)
            validation.add(f"{column}4:{column}500")
        for column in ("F", "L"):
            for cell in sheet[column][3:]:
                cell.number_format = "dd/mm/yyyy"

        context = workbook.create_sheet("Contexto")
        context.append(["Cliente", "NIT / ERP", "Asesor", "Última compra", "Ventas otras marcas", "Marcas compradas", "Marcas lineales detectadas"])
        for customer in customers:
            context.append([
                customer["customer_name"], customer["customer_erp_id"], customer["advisor"],
                customer.get("last_purchase"), float(customer.get("other_sales") or 0),
                customer.get("brands") or "", customer.get("competitor_linear_brands") or "Ninguna detectada",
            ])
        for cell in context[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=dark_blue)
        context.freeze_panes = "A2"
        context.auto_filter.ref = context.dimensions
        context.column_dimensions["A"].width = 36
        context.column_dimensions["B"].width = 16
        context.column_dimensions["C"].width = 24
        context.column_dimensions["D"].width = 16
        context.column_dimensions["E"].width = 22
        context.column_dimensions["F"].width = 55
        context.column_dimensions["G"].width = 28
        for cell in context["E"][1:]:
            cell.number_format = '"COP "#,##0'

        guide = workbook.create_sheet("Instrucciones")
        guide.append(["Campaña THK para clientes de plásticos"])
        guide.append(["Objetivo", "Identificar aplicaciones de movimiento lineal y abrir oportunidades de homologación THK."])
        guide.append(["Primera actualización", "27/10/2026"])
        guide.append(["Segundo y último plazo", "02/11/2026"])
        guide.append(["Escalamiento a Gerencia", "A partir del 04/11/2026 si no se registra ninguna actualización."])
        guide.append(["Qué debe registrar", "Contacto, aplicación o máquina, referencia instalada, interés y próxima acción."])
        guide.append(["Qué ocurre después", "Los clientes marcados como oportunidad pasan al proceso comercial normal del Command Center."])
        guide["A1"].font = Font(size=18, bold=True, color="FFFFFF")
        guide["A1"].fill = PatternFill("solid", fgColor=dark_blue)
        guide.merge_cells("A1:B1")
        for cell in guide["A"][1:]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=orange)
        guide.column_dimensions["A"].width = 28
        guide.column_dimensions["B"].width = 90
        for row in guide.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        safe_advisor = (advisor or customers[0]["advisor"]).title().replace(" ", "_")
        return stream, f"Campana_THK_Plasticos_{safe_advisor}.xlsx"

    @staticmethod
    @transactional
    def import_workbook(initiative_id: int, file_stream) -> dict[str, int]:
        workbook = load_workbook(file_stream, data_only=False)
        if "Seguimiento" not in workbook.sheetnames:
            raise ValueError("El archivo no contiene la hoja Seguimiento.")
        sheet = workbook["Seguimiento"]
        headers = [cell.value for cell in sheet[3]]
        if headers[: len(TRACKING_HEADERS)] != TRACKING_HEADERS:
            raise ValueError("La estructura del archivo cambió. Descargue una copia nueva.")
        updated = 0
        ignored = 0
        for row in sheet.iter_rows(min_row=4, values_only=True):
            source_key = str(row[0] or "").strip()
            if not source_key:
                continue
            values = {
                "status": row[4], "contact_date": InitiativeCampaignService._date(row[5]),
                "contact_name": row[6], "application": row[7],
                "installed_reference": row[8], "interest": row[9],
                "next_action": row[10], "next_action_date": InitiativeCampaignService._date(row[11]),
                "opportunity_decision": row[12],
            }
            if InitiativeCampaignRepository.update_tracking(initiative_id, source_key, values):
                updated += 1
            else:
                ignored += 1
        InitiativeRepository.create_event(
            initiative_id=initiative_id,
            event_type="campaign_sync",
            title="Seguimiento de campaña sincronizado",
            details=f"{updated} clientes actualizados; {ignored} filas ignoradas.",
            created_by="excel",
        )
        return {"updated": updated, "ignored": ignored}

    @staticmethod
    def _date(value):
        if value is None or value == "":
            return None
        if hasattr(value, "date"):
            return value.date().isoformat()
        return str(value).strip()
